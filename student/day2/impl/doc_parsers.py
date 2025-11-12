# -*- coding: utf-8 -*-
from __future__ import annotations
from pathlib import Path
from typing import List, Dict, Any
import re, json, csv

# ───────────── 기본 유틸 ─────────────
def _normalize_text(s: str) -> str:
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()

# ───────────── 포맷별 리더 ─────────────
def _read_pdf(p: Path) -> str:
    from pdfminer.high_level import extract_text
    return extract_text(str(p)) or ""

def _read_docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    return "\n".join(para.text for para in doc.paragraphs)

def _read_xlsx(p: Path) -> tuple[str, List[List[str]]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), data_only=True)
    rows: List[List[str]] = []
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            rows.append([("" if c is None else str(c)).strip() for c in r])
    # 텍스트로도 합쳐서 추출기로 넘김
    text = "\n".join(" | ".join(r) for r in rows)
    return text, rows

def _read_csv(p: Path) -> str:
    # UTF-8 기본, 실패 시 cp949 백업
    for enc in ("utf-8", "cp949"):
        try:
            with open(p, newline="", encoding=enc) as f:
                rd = csv.reader(f)
                return "\n".join(" | ".join(c.strip() for c in row) for row in rd)
        except Exception:
            continue
    # 최후: 바이너리 읽고 디코딩 무시
    with open(p, "rb") as f:
        return f.read().decode("utf-8", errors="ignore")

def _read_md(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="ignore")

def _read_json(p: Path) -> str:
    def flatten(o) -> str:
        if isinstance(o, dict):
            return " ".join(f"{k}: {flatten(v)}" for k, v in o.items())
        if isinstance(o, list):
            return " ".join(flatten(x) for x in o)
        return str(o)
    obj = json.loads(p.read_text(encoding="utf-8", errors="ignore"))
    return flatten(obj)

def _read_txt(p: Path) -> str:
    for enc in ("utf-8", "cp949"):
        try:
            return p.read_text(encoding=enc)
        except Exception:
            continue
    return p.read_text(encoding="utf-8", errors="ignore")

# ───────────── 공개 함수 ─────────────
def load_any(path: str) -> Dict[str, Any]:
    p = Path(path)
    ext = p.suffix.lower()  # 대소문자 무시
    if ext == ".pdf":
        txt = _read_pdf(p)
        return {"type": "pdf", "text": _normalize_text(txt), "tables": []}

    if ext in (".docx", ".doc"):
        txt = _read_docx(p)
        return {"type": "docx", "text": _normalize_text(txt), "tables": []}

    if ext in (".xlsx", ".xls"):
        text, rows = _read_xlsx(p)
        return {"type": "xlsx", "text": _normalize_text(text), "tables": rows}

    if ext == ".csv":
        txt = _read_csv(p)
        return {"type": "csv", "text": _normalize_text(txt), "tables": []}

    if ext == ".md":
        txt = _read_md(p)
        return {"type": "md", "text": _normalize_text(txt), "tables": []}

    if ext == ".json":
        txt = _read_json(p)
        return {"type": "json", "text": _normalize_text(txt), "tables": []}

    # 기본: txt
    return {"type": "txt", "text": _normalize_text(_read_txt(p)), "tables": []}
